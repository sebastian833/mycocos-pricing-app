#!/usr/bin/env python3
"""
Revisión de costos: toma el costo NETO de la maestra (pestaña vigente por marca) y lo compara
con lo que usa hoy la app. Sin argumentos genera solo el reporte (no modifica nada).
Con --aplicar actualiza los JSON del repo en PUB.

Pestañas usadas (costo neto):
  MyCOCOS            -> 'MyCOCOS'
  MyHUEVOS Colombia  -> 'MYHUEVOS COL Antigua'  (confirmado por el usuario)
  MENNT              -> 'E COMPANY'
"""
import openpyxl, json, re, sys, os, unicodedata, warnings
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

warnings.filterwarnings('ignore')

import argparse
_ap = argparse.ArgumentParser(description='Alinea los costos de la app con la maestra (costo neto por SKU).')
_ap.add_argument('--aplicar', action='store_true', help='escribe los cambios en public/*.json (sin esto solo genera el reporte)')
_ap.add_argument('--repo', default='.', help='carpeta del proyecto (la que tiene public/)')
_ap.add_argument('--datos', default='../datos-fuente', help='carpeta datos-fuente')
_ap.add_argument('--maestra', help='ruta de la maestra (.xlsx); por defecto datos-fuente/maestra/MAESTRA_PRODUCTOS.xlsx')
_ap.add_argument('--reporte', default='../reportes/reporte_cambios_costos.xlsx', help='dónde dejar el reporte Excel')
_args = _ap.parse_args()
APLICAR = _args.aplicar
PUB = os.path.join(_args.repo, 'public')
UP = os.path.join(_args.datos, 'ventas')
OUT = _args.reporte
CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.cache', 'erp_names.json')
MAESTRA = _args.maestra or os.path.join(_args.datos, 'maestra', 'MAESTRA_PRODUCTOS.xlsx')
IVA = 1.19
TOL = 0.005  # 0,5 %

wb = openpyxl.load_workbook(MAESTRA, data_only=True)


def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9]+', ' ', s)).strip()


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


# ---------------------------------------------------------------- maestra
def maestra_mycocos():
    ws, out = wb['MyCOCOS'], {}
    for row in ws.iter_rows(min_row=1, max_row=130):
        sku = row[0].value
        if isinstance(sku, str) and sku.startswith('MC-'):
            f = num(row[5].value)
            if f and f > 100:
                out.setdefault(sku, dict(nombre=str(row[1].value).strip(), neto=round(f), estado=''))
    return out


def maestra_myhuevos():
    ws, out = wb['MYHUEVOS COL Antigua'], {}
    for row in ws.iter_rows(min_row=1):
        sku = row[0].value
        if isinstance(sku, str) and sku.startswith('MH'):
            e = num(row[4].value)
            if e and e > 100:
                out.setdefault(sku, dict(nombre=str(row[1].value).strip(), neto=round(e), estado=''))
    return out


def maestra_mennt():
    ws, out = wb['E COMPANY'], {}
    for row in ws.iter_rows(min_row=1):
        sku = row[0].value
        if isinstance(sku, str) and sku.startswith('MN-'):
            h = num(row[7].value)
            if h and h > 100:
                out.setdefault(sku, dict(nombre=str(row[3].value).strip(), neto=round(h),
                                         estado=str(row[4].value or '').strip().upper()))
    return out


MAES = {'mycocos': maestra_mycocos(), 'myhuevos': maestra_myhuevos(), 'mennt': maestra_mennt()}
NOMBRE_MARCA = {'mycocos': 'MyCOCOS', 'myhuevos': 'MyHUEVOS Colombia', 'mennt': 'MENNT'}


# ---------------------------------------------------------------- nombre -> SKU (reporte de ventas)
def cargar_erp():
    res = json.load(open(CACHE)) if os.path.exists(CACHE) else {}
    rutas = {
        'mycocos': [f'{UP}/mycocos_ventas_2025.xlsx',
                    f'{UP}/mycocos_ventas_2026_hasta-17-jul.xlsx'],
        'mennt': [f'{UP}/mennt_ventas_2025.xlsx',
                  f'{UP}/mennt_ventas_2026_hasta-18-sep.xlsx'],
    }
    cambiado = False
    for marca, paths in rutas.items():
        if marca in res:
            continue
        m = {}
        for p in paths:
            df = pd.read_excel(p, usecols=['Tipo Movimiento', 'Producto / Servicio', 'SKU'])
            df = df[df['Tipo Movimiento'] == 'venta']
            for name, sku in df.groupby('Producto / Servicio')['SKU'].agg(lambda x: x.mode()[0]).items():
                m.setdefault(str(name), str(sku))
        res[marca] = m
        cambiado = True
    if 'myhuevos' not in res:  # Colombia: el reporte trae el SKU de la maestra en la columna SKU/EAN
        m = {}
        for p in [f'{UP}/myhuevos-col_ventas_2026_hasta-31-ago.xlsx',
                  f'{UP}/myhuevos-col_ventas_2025.xlsx']:
            df = pd.read_excel(p, usecols=['Producto', 'SKU/EAN', 'Estado de la Orden'])
            df = df[df['Estado de la Orden'] == 'Entregado']
            for name, sku in df.groupby('Producto')['SKU/EAN'].agg(lambda x: str(x.mode()[0])).items():
                m.setdefault(str(name), str(sku))
        res['myhuevos'] = m
        cambiado = True
    if cambiado:
        os.makedirs(os.path.dirname(CACHE), exist_ok=True)
        json.dump(res, open(CACHE, 'w'), ensure_ascii=False)
    return res


ERP = cargar_erp()


def sku_por_nombre(marca, nombre):
    maes = MAES[marca]
    m = ERP.get(marca, {})
    sku = m.get(nombre)
    if not sku:
        sku = next((v for k, v in m.items() if k[:70] == nombre), None)
    if sku in maes:
        return sku
    idx = {norm(v['nombre']): k for k, v in maes.items()}
    return idx.get(norm(nombre))


def cargar(nombre):
    return json.load(open(f'{PUB}/{nombre}'))


def guardar(nombre, data, indent=None):
    with open(f'{PUB}/{nombre}', 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=indent)


# ---------------------------------------------------------------- 1. productos (Simulador / Informe)
filas_prod, datos = [], {}
for marca in MAES:
    d = cargar(f'data-{marca}.json')
    datos[marca] = d
    for p in d['productos']:
        sku = sku_por_nombre(marca, p['nombre'])
        neto = MAES[marca][sku]['neto'] if sku else None
        nota = '' if sku else 'No encontré este producto en la maestra (se queda como está)'
        filas_prod.append(dict(marca=marca, nombre=p['nombre'], sku=sku or '', actual=p['costo_avg'], neto=neto, nota=nota))

# ---------------------------------------------------------------- 2. catálogo del armador / regalos
filas_cat, cats = [], {}
for marca in MAES:
    cat = cargar(f'componentes-{marca}.json')
    cats[marca] = cat
    for it in cat:
        sku = it['sku'] if it['sku'] in MAES[marca] else sku_por_nombre(marca, it['nombre'])
        neto = MAES[marca][sku]['neto'] if sku else None
        estado = MAES[marca][sku]['estado'] if sku else ''
        nota = ''
        if neto is None:
            nota = 'No encontré este producto en la maestra (se queda como está)'
        elif abs(it['costo'] - neto * IVA) <= 3:
            nota = 'El costo actual es el bruto (con IVA)'
        elif estado == 'INACTIVO':
            nota = 'Inactivo en la maestra'
        filas_cat.append(dict(marca=marca, sku=sku or it['sku'], nombre=it['nombre'], actual=it['costo'],
                              neto=neto, estado=estado, nota=nota))

# ---------------------------------------------------------------- 3. packs
filas_pack = []

# 3a. MyCOCOS: packs ganadores (por componente)
REGLAS = [
    ('4 0', ['rasuradora', '4 0', 'black'], ['kit', 'hojilla']),
    ('daily shaver', ['daily shaver'], ['hojilla']),
    ('nariz y orejas', ['nariz y orejas'], ['cabezal']),
    ('safety razor', ['safety razor'], ['hojilla']),
    ('travel bag', ['travel bag', '1 0'], []),
    ('cuchilla de repuesto', ['5 hojillas', 'safety'], []),
    ('desodorante', ['desodorante', '30'], []),
    ('cortau', ['cortau', 'pro'], []),
    ('espuma', ['espuma'], []),
    ('balsamo', ['balsamo'], []),
    ('wipe', ['wipes to go'], []),
    ('jabon', ['pack 3 jabones'], []),
    ('cepillo', ['cepillo'], []),
]


def hallar(maes, incl, excl):
    for sku, v in maes.items():
        n = norm(v['nombre'])
        if all(norm(t) in n for t in incl) and not any(norm(t) in n for t in excl):
            return sku
    return None


def costo_componente_mycocos(nombre):
    n = norm(nombre)
    if 'cepillo' in n:  # el regalo dice "cepillos" (plural): capilar + corporal
        a = hallar(MAES['mycocos'], ['cepillo', 'capilar'], [])
        b = hallar(MAES['mycocos'], ['cepillo', 'corporal'], [])
        if a and b:
            return (MAES['mycocos'][a]['neto'] + MAES['mycocos'][b]['neto'], a)
    for clave, incl, excl in REGLAS:
        if clave in n:
            sku = hallar(MAES['mycocos'], incl, excl)
            return (MAES['mycocos'][sku]['neto'], sku) if sku else (None, None)
    return (None, None)


gan_my = cargar('packs-ganadores-mycocos.json')
for pk in gan_my:
    viejo = nuevo = 0
    sin = []
    for c in pk['componentes']:
        costo_v = c['costo'] or 0
        nuevo_c, _ = costo_componente_mycocos(c['nombre'])
        if nuevo_c is None:
            nuevo_c = costo_v
            sin.append(c['nombre'])
        viejo += costo_v * c['qty']
        nuevo += nuevo_c * c['qty']
    filas_pack.append(dict(marca='mycocos', pack=pk['nombre'] + ' (ganador)', actual=viejo, nuevo=nuevo,
                           nota=('Sin costo en la maestra, se mantiene: ' + ', '.join(sin)) if sin else ''))

# 3b. MyCOCOS: recetas (packs.json)
packs_my = cargar('packs-mycocos.json')
for pk in packs_my:
    if not (pk.get('precio_venta') and pk['costo_real_bom'] > 0):
        continue
    nuevo, sin = 0, 0
    for c in pk['componentes']:
        m = MAES['mycocos'].get(c['sku'])
        unit = m['neto'] if m else c['costo_unit']
        if not m:
            sin += 1
        nuevo += unit * c['qty']
    filas_pack.append(dict(marca='mycocos', pack=pk['nombre'], actual=pk['costo_real_bom'], nuevo=round(nuevo),
                           nota=f'{sin} componente(s) sin costo en la maestra, se mantienen' if sin else ''))

# 3c. MyHUEVOS Colombia: packs ganadores (solo candidatos, no se cambia solo)
gan_mh = cargar('packs-ganadores-myhuevos.json')
CLAVES_MH = [('kit prime', ['prime']), ('esencial', ['esencial']), ('master pack', ['master']),
             ('perfect couple', ['perfect', 'couple'])]
for pk in gan_mh:
    n = norm(pk['nombre'])
    actual = sum((c['costo'] or 0) * c['qty'] for c in pk['componentes'])
    cands = []
    for clave, toks in CLAVES_MH:
        if clave in n:
            color = 'negro' if 'negro' in n else ('azul' if 'azul' in n else None)
            for sku, v in MAES['myhuevos'].items():
                vn = norm(v['nombre'])
                if all(t in vn for t in toks) and (color is None or color in vn):
                    cands.append(f"{v['nombre'][:45]} = ${v['neto']:,}")
    filas_pack.append(dict(marca='myhuevos', pack=pk['nombre'] + ' (ganador)', actual=actual, nuevo=None,
                           nota=('Revisar, candidatos en la maestra: ' + ' | '.join(cands[:3])) if cands
                           else 'No encontré este pack en la pestaña Antigua, revisar'))

# ---------------------------------------------------------------- resumen en consola
def cambia(a, b):
    return b is not None and a and abs(b / a - 1) >= TOL


def cambia5(a, b):
    return b is not None and a and abs(b / a - 1) >= 0.05


print("=== RESUMEN (reporte, aún no se aplicó nada)" if not APLICAR else "=== APLICANDO CAMBIOS")
for marca in MAES:
    fp = [f for f in filas_prod if f['marca'] == marca]
    fc = [f for f in filas_cat if f['marca'] == marca]
    fk = [f for f in filas_pack if f['marca'] == marca]
    print(f"\n{NOMBRE_MARCA[marca]}")
    print(f"  Simulador/Informe: {len(fp)} productos | cambian {sum(cambia(f['actual'], f['neto']) for f in fp)} (más de 5%: {sum(cambia5(f['actual'], f['neto']) for f in fp)}) | sin match {sum(f['neto'] is None for f in fp)}")
    print(f"  Catálogo armador:  {len(fc)} productos | cambian {sum(cambia(f['actual'], f['neto']) for f in fc)} (más de 5%: {sum(cambia5(f['actual'], f['neto']) for f in fc)}) | "
          f"con IVA incluido {sum('bruto' in f['nota'] for f in fc)} | sin match {sum(f['neto'] is None for f in fc)} | inactivos {sum(f['estado']=='INACTIVO' for f in fc)}")
    print(f"  Packs:             {len(fk)} | cambian {sum(cambia(f['actual'], f['nuevo']) for f in fk)} (más de 5%: {sum(cambia5(f['actual'], f['nuevo']) for f in fk)}) | a revisar {sum(f['nuevo'] is None for f in fk)}")

# ---------------------------------------------------------------- reporte Excel
ARIAL = 'Arial'
f_norm = Font(name=ARIAL, size=10)
f_head = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
fill_head = PatternFill('solid', start_color='1F2937')
fill_input = PatternFill('solid', start_color='FFF2CC')

xl = Workbook()


def hoja(titulo, headers, anchos, filas, fmt_cols):
    ws = xl.create_sheet(titulo)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = f_head, fill_head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for fila in filas:
        ws.append(fila)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = f_norm
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[L(i)].width = w
    for col, fmt in fmt_cols.items():
        for r in range(2, ws.max_row + 1):
            ws[f'{col}{r}'].number_format = fmt
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


# Productos: A Marca B Producto C SKU D Actual E Maestra F Dif$ G Dif% H Estado I Nota
filas = []
for i, f in enumerate(sorted(filas_prod, key=lambda x: (x['marca'], x['nombre'])), start=2):
    filas.append([NOMBRE_MARCA[f['marca']], f['nombre'], f['sku'], f['actual'], f['neto'],
                  f'=IF(E{i}="","",E{i}-D{i})', f'=IF(OR(E{i}="",D{i}=0),"",E{i}/D{i}-1)',
                  f'=IF(E{i}="","Sin match",IF(ABS(G{i})<{TOL},"Igual","Cambia"))', f['nota']])
hoja('Productos', ['Marca', 'Producto (Simulador / Informe)', 'SKU', 'Costo actual en la app', 'Costo neto maestra',
                   'Diferencia $', 'Diferencia %', 'Estado', 'Nota'],
     [18, 52, 20, 16, 16, 14, 13, 12, 60], filas, {'D': '#,##0', 'E': '#,##0', 'F': '#,##0;-#,##0;-', 'G': '0.0%'})

# Catálogo: A Marca B SKU C Producto D Actual E Maestra F Dif$ G Dif% H Estado I EstadoMaestra J Nota
filas = []
for i, f in enumerate(sorted(filas_cat, key=lambda x: (x['marca'], x['nombre'])), start=2):
    filas.append([NOMBRE_MARCA[f['marca']], f['sku'], f['nombre'], f['actual'], f['neto'],
                  f'=IF(E{i}="","",E{i}-D{i})', f'=IF(OR(E{i}="",D{i}=0),"",E{i}/D{i}-1)',
                  f'=IF(E{i}="","Sin match",IF(ABS(G{i})<{TOL},"Igual","Cambia"))', f['estado'], f['nota']])
hoja('Catálogo armador', ['Marca', 'SKU', 'Producto', 'Costo actual en la app', 'Costo neto maestra',
                          'Diferencia $', 'Diferencia %', 'Estado', 'Estado en maestra', 'Nota'],
     [18, 20, 52, 16, 16, 14, 13, 12, 14, 52], filas, {'D': '#,##0', 'E': '#,##0', 'F': '#,##0;-#,##0;-', 'G': '0.0%'})

# Packs: A Marca B Pack C Actual D Nuevo E Dif$ F Dif% G Estado H Nota
filas = []
for i, f in enumerate(sorted(filas_pack, key=lambda x: (x['marca'], x['pack'])), start=2):
    filas.append([NOMBRE_MARCA[f['marca']], f['pack'], f['actual'], f['nuevo'],
                  f'=IF(D{i}="","",D{i}-C{i})', f'=IF(OR(D{i}="",C{i}=0),"",D{i}/C{i}-1)',
                  f'=IF(D{i}="","Revisar",IF(ABS(F{i})<{TOL},"Igual","Cambia"))', f['nota']])
hoja('Packs', ['Marca', 'Pack', 'Costo actual en la app', 'Costo nuevo (maestra)', 'Diferencia $', 'Diferencia %',
               'Estado', 'Nota'],
     [18, 62, 16, 16, 14, 13, 12, 80], filas, {'C': '#,##0', 'D': '#,##0', 'E': '#,##0;-#,##0;-', 'F': '0.0%'})

# Resumen con fórmulas
ws = xl.active
ws.title = 'Resumen'
ws['A1'] = 'Revisión de costos: maestra (neto) vs. lo que usa hoy la app'
ws['A1'].font = Font(name=ARIAL, size=13, bold=True)
enc = ['Marca', 'Productos que cambian', '...de ellos, más de 5 %', 'Productos sin match',
       'Catálogo: cambian', '...de ellos, más de 5 %', 'Catálogo: hoy con IVA incluido', 'Catálogo: sin match',
       'Packs que cambian', '...de ellos, más de 5 %', 'Packs a revisar']
for j, h in enumerate(enc, 1):
    c = ws.cell(row=3, column=j, value=h)
    c.font, c.fill = f_head, fill_head
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
CAT = "'Catálogo armador'"
for r, marca in enumerate(NOMBRE_MARCA.values(), start=4):
    ws.cell(row=r, column=1, value=marca)
    ws.cell(row=r, column=2, value=f'=COUNTIFS(Productos!A:A,$A{r},Productos!H:H,"Cambia")')
    ws.cell(row=r, column=3, value=f'=COUNTIFS(Productos!A:A,$A{r},Productos!G:G,">0.05")+COUNTIFS(Productos!A:A,$A{r},Productos!G:G,"<-0.05")')
    ws.cell(row=r, column=4, value=f'=COUNTIFS(Productos!A:A,$A{r},Productos!H:H,"Sin match")')
    ws.cell(row=r, column=5, value=f'=COUNTIFS({CAT}!A:A,$A{r},{CAT}!H:H,"Cambia")')
    ws.cell(row=r, column=6, value=f'=COUNTIFS({CAT}!A:A,$A{r},{CAT}!G:G,">0.05")+COUNTIFS({CAT}!A:A,$A{r},{CAT}!G:G,"<-0.05")')
    ws.cell(row=r, column=7, value=f'=COUNTIFS({CAT}!A:A,$A{r},{CAT}!J:J,"*bruto*")')
    ws.cell(row=r, column=8, value=f'=COUNTIFS({CAT}!A:A,$A{r},{CAT}!H:H,"Sin match")')
    ws.cell(row=r, column=9, value=f'=COUNTIFS(Packs!A:A,$A{r},Packs!G:G,"Cambia")')
    ws.cell(row=r, column=10, value=f'=COUNTIFS(Packs!A:A,$A{r},Packs!F:F,">0.05")+COUNTIFS(Packs!A:A,$A{r},Packs!F:F,"<-0.05")')
    ws.cell(row=r, column=11, value=f'=COUNTIFS(Packs!A:A,$A{r},Packs!G:G,"Revisar")')
for row in ws.iter_rows(min_row=4, max_row=6):
    for c in row:
        c.font = f_norm
notas = [
    'Reglas y supuestos',
    'Costo = costo NETO (sin IVA) de la maestra, por SKU. Pestañas: MyCOCOS -> "MyCOCOS"; MyHUEVOS Colombia -> "MYHUEVOS COL Antigua" (indicado por ti); MENNT -> "E COMPANY".',
    'Diferencias menores a 0,5 % se consideran "Igual". IVA 19 % para detectar costos que hoy están en bruto.',
    'No se tocan los costos históricos mes a mes de MyCOCOS y MENNT en el Informe (vienen del reporte de ventas y son lo que realmente se pagó en cada época).',
    'Los packs ganadores de MyHUEVOS no se cambian solos: salen como "Revisar" con los candidatos de la maestra.',
    'Productos o componentes sin match se quedan con el costo actual; revísalos en las columnas Nota.',
]
for k, t in enumerate(notas):
    c = ws.cell(row=9 + k, column=1, value=t)
    c.font = Font(name=ARIAL, size=10, bold=(k == 0))
for col, w in zip('ABCDEFGHIJK', [22, 14, 14, 14, 14, 14, 18, 14, 14, 14, 14]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[3].height = 42

os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
xl.save(OUT)
print(f"\nReporte guardado: {OUT}")

# ---------------------------------------------------------------- aplicar
if APLICAR:
    # data-*.json
    for marca, d in datos.items():
        for p in d['productos']:
            sku = sku_por_nombre(marca, p['nombre'])
            if not sku:
                continue
            neto = MAES[marca][sku]['neto']
            if not cambia(p['costo_avg'], neto):
                continue
            p['costo_avg'] = neto
            pn = p['precio_bruto_avg'] / IVA
            p['margen_avg'] = round((pn - neto) / pn * 100, 1) if pn > 0 else 0
            if marca == 'myhuevos':  # sin costo en el reporte de ventas: costo plano mes a mes
                for m in p['meses']:
                    m['c25'] = neto
                    m['mg25'] = round((m['pn25'] - neto) / m['pn25'] * 100, 1) if m['pn25'] > 0 else 0
        guardar(f'data-{marca}.json', d)
    # componentes
    for marca, cat in cats.items():
        for it in cat:
            sku = it['sku'] if it['sku'] in MAES[marca] else sku_por_nombre(marca, it['nombre'])
            if sku:
                it['costo'] = MAES[marca][sku]['neto']
        guardar(f'componentes-{marca}.json', cat)
    # packs ganadores MyCOCOS
    for pk in gan_my:
        for c in pk['componentes']:
            nuevo_c, _ = costo_componente_mycocos(c['nombre'])
            if nuevo_c is not None:
                c['costo'] = nuevo_c
    guardar('packs-ganadores-mycocos.json', gan_my, indent=2)
    # recetas MyCOCOS
    for pk in packs_my:
        total = 0
        for c in pk['componentes']:
            m = MAES['mycocos'].get(c['sku'])
            if m:
                c['costo_unit'] = m['neto']
            total += c['costo_unit'] * c['qty']
        if pk.get('precio_venta') and pk['costo_real_bom'] > 0:
            pk['costo_real_bom'] = round(total)
    guardar('packs-mycocos.json', packs_my)
    print("Archivos del repo actualizados.")
